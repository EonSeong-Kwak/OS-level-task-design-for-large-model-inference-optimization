#!/usr/bin/env python3
"""Generate the full technical design document as .docx."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "面向大模型推理优化的操作系统级任务_完整技术方案.docx"


def set_run_font(run, name="宋体", size=12, bold=False, color=None, west="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = west
    if color is not None:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml import OxmlElement

        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), west)
    rFonts.set(qn("w:hAnsi"), west)
    rFonts.set(qn("w:eastAsia"), name)


def add_heading_cn(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "黑体", 16 if level == 1 else 14 if level == 2 else 12, bold=True)
    return p


def add_p(doc, text, *, size=12, first_line=True, bold=False, align="justify"):
    p = doc.add_paragraph()
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    if first_line:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, "宋体", size, bold=bold)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Cm(0.75 + 0.75 * level)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, "宋体", 12)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, "宋体", 9, west="Consolas")
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "1F4E79")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, "黑体", 10.5, bold=True, west="Calibri")
        shade_header(hdr[i])
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, "宋体", 10.5, west="Calibri")
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)

    # cover
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("计算机科学与技术专业实习")
    set_run_font(r, "黑体", 18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A 类选题 · 系统设计方案")
    set_run_font(r, "黑体", 16, bold=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向大模型推理优化的操作系统级任务")
    set_run_font(r, "黑体", 22, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("完整技术方案")
    set_run_font(r, "黑体", 20, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("llmOS：可仿真、可评测、可对照真实引擎的推理操作系统内核")
    set_run_font(r, "楷体", 14)

    doc.add_paragraph()
    meta = [
        "适用对象：高年级本科 / 研究生实习或《操作系统进阶》《大模型系统》实训",
        "分组规模：4 人（架构 / 内存与系统 / 运行时与 I/O / 测试评测）",
        "周期：12 周 / 96 学时（可按实习节点压缩为设计 2 周 + 实现 5 周 + 测试 1–2 周）",
        "硬件：本地 Python 仿真为主；云端 GPU Notebook 仅作真实性背书",
        "对应课程文档：面向大模型推理优化的操作系统级任务设计（设计 + 原理解释 合并版）",
        "版本：v1.0  |  日期：2026-09-14",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, "宋体", 12)

    doc.add_page_break()

    add_heading_cn(doc, "1  项目背景与目标", 1)
    add_heading_cn(doc, "1.1  问题陈述", 2)
    add_p(
        doc,
        "大模型推理服务在运行时呈现出与传统操作系统高度同构的资源矛盾：大量并发请求竞争有限的 GPU 算力与显存；"
        "KV Cache 随自回归生成只增不减且长度事先未知；超大模型权重无法一次性装入单卡；多租户共享时既要隔离又要提高承载密度。"
        "工业界的 vLLM PagedAttention、Continuous Batching、权重卸载与 GPU 虚拟化（MPS/MIG），本质上是把操作系统中已经成熟的分页、调度、异步 I/O 与超卖思想，迁移到了推理这条新的 workload 上。",
    )
    add_p(
        doc,
        "本方案不把任务做成互不相关的四份作业，而是建设一套名为 llmOS 的推理操作系统内核：用同一套离散事件时钟驱动内存管理、进程调度、存储栈与虚拟化四个子系统，使一次推理请求能够穿过完整的 OS 路径，并在统一指标总线上给出可复现的量化对比。90% 工作量在本地仿真器完成，仅在最后 1–2 周用 Colab / Kaggle / AutoDL 对 1–2 个关键机制做真实性背书，满足课程“零本地硬件、可规模化排课”的约束。",
    )

    add_heading_cn(doc, "1.2  建设目标", 2)
    add_p(doc, "功能性目标：")
    add_bullet(doc, "实现 KV Cache 分页内存管理：块分配、页表翻译、按需扩页、写时复制、LRU 换出。")
    add_bullet(doc, "实现持续批处理调度器：静态批基线、迭代级组批、FCFS/SJF/抢占/公平队列，并区分 Prefill 与 Decode。")
    add_bullet(doc, "实现权重预取与计算-I/O 重叠运行时，以及类 cgroups 的算力份额与显存硬上限。")
    add_bullet(doc, "实现 GPU 虚拟地址空间、时间片复用、显存超卖、balloon 回收与越界拦截。")
    add_p(doc, "非功能性目标：")
    add_bullet(doc, "仿真层纯 Python 3.10+，一键复现；单测覆盖率不低于 70%。")
    add_bullet(doc, "所有结论必须来自同一 trace、同一随机种子下的基线 vs 优化对比，禁止只展示优化后的绝对数。")
    add_bullet(doc, "四个子系统通过冻结的接口契约集成，角色 B/C 可并行开发、强制互审。")
    add_bullet(doc, "真实 GPU 实验失败不影响主交付：仿真结果本身构成可答辩系统。")

    add_heading_cn(doc, "1.3  非目标（明确不做）", 2)
    add_bullet(doc, "不从零实现 CUDA Attention kernel，不维护完整 vLLM 分支。")
    add_bullet(doc, "不以刷榜为目的调参；不以不可复现的云端环境作为唯一证据。")
    add_bullet(doc, "不把四个任务做成四套互不相通的全局变量脚本。")

    add_heading_cn(doc, "1.4  成功判据", 2)
    add_table(
        doc,
        ["维度", "最低可答辩标准", "加分标准"],
        [
            [
                "系统联通",
                "一条请求能依次经过调度 → 配额 → VMM → KV 分页",
                "多租户 + 超卖 + 换出同时在线",
            ],
            [
                "KV 内存",
                "分页相对连续预占：碎片下降、利用率上升、OOM 减少",
                "COW 前缀共享 + LRU 换出有独立消融",
            ],
            [
                "调度",
                "持续批相对静态批：吞吐上升、TTFT 改善",
                "FAIR 使多租户 P99 方差下降且饥饿占比可控",
            ],
            [
                "I/O",
                "双缓冲相对串行：I/O stall 占比下降",
                "lookahead 与页缓存有消融曲线",
            ],
            [
                "虚拟化",
                "逻辑页 > 物理页可运行；越界拦截率 = 0",
                "超卖比–延迟 Pareto 曲线 + Jain 公平指数",
            ],
            [
                "工程",
                "pytest 覆盖率 ≥ 70%，README 一键启动",
                "云端一项对照实验与仿真趋势一致",
            ],
        ],
    )

    add_heading_cn(doc, "2  操作系统原理映射", 1)
    add_p(
        doc,
        "推理服务是一个多道程序系统：请求类比进程，KV Cache 与权重类比地址空间中的数据与代码，GPU 类比 CPU，NVMe/CPU/HBM 类比存储层次。四个任务分别对应 OS 四大子系统，本方案在实现层保持这一映射，避免滑向“纯算法调参”。",
    )
    add_table(
        doc,
        ["任务", "OS 子系统", "核心矛盾", "经典机制（必须实现）"],
        [
            [
                "一、KV 分页",
                "虚拟内存",
                "显存碎片与按最大值预占浪费",
                "分页、页表、demand paging、COW、置换、块分配器",
            ],
            [
                "二、持续批处理",
                "进程调度",
                "GPU 空闲与尾延迟、饥饿",
                "PCB 状态机、FCFS/SJF/RR/MLFQ、抢占、公平队列",
            ],
            [
                "三、预取 + 配额",
                "I/O 与资源管理",
                "搬运 stall 与租户争用",
                "异步 I/O、双缓冲、预取、页缓存、cgroups 配额",
            ],
            [
                "四、GPU 虚拟化",
                "虚拟化",
                "单卡租户数 vs 隔离",
                "GPA/HPA、超卖、页错误、balloon、时间片、越界防护",
            ],
        ],
    )
    add_p(
        doc,
        "量化直觉（作为设计预期，最终以实验为准）：分页使 KV 从“按 max_tokens 预占”变为“按实际生成占用”，工业界 PagedAttention 可将碎片压到接近 0，并提升单卡并发；持续批处理消除静态批中短请求被长请求拖死造成的空泡；计算-I/O 重叠可将 stall 从 30% 量级压到个位数；超卖用时间换空间，单卡逻辑承载可到 1.5×–3×，代价是换入换出延迟。",
    )

    add_heading_cn(doc, "3  总体架构", 1)
    add_heading_cn(doc, "3.1  分层与数据流", 2)
    add_p(
        doc,
        "llmOS 采用单时钟、多子系统的微内核式仿真：所有时间推进由 DiscreteEventSim 独占。上层只投递事件，下层只响应接口调用。请求对象 Request 作为贯穿全栈的 PCB，其 rid、tenant、priority、prompt_len、max_tokens 以及生成进度由调度器维护；物理资源记账由 VMM 与 KV 分配器维护。",
    )
    add_code(
        doc,
        "Workload(Trace/Poisson)\n"
        "        |  Arrival 事件\n"
        "        v\n"
        " DiscreteEventSim  (now, 优先队列, schedule/run/step)\n"
        "        |\n"
        "        +--> Scheduler._pick_batch()     # 任务二：决定本 iteration 集合\n"
        "        |         |\n"
        "        |         +--> QuotaManager       # 任务三：算力份额 / 显存硬上限\n"
        "        |         +--> MPSScheduler       # 任务四：时间片与 ctx-switch 开销\n"
        "        |         +--> PrefetchRuntime    # 任务三：层间搬运与 stall 记账\n"
        "        |         v\n"
        "        +--> PageTable.append_token()    # 任务一：逻辑 token -> KV block\n"
        "                  |\n"
        "                  +--> BlockAllocator     # 引用计数 / COW / LRU\n"
        "                  +--> GPUVMM.access()    # 任务四：驻留、换出、隔离\n"
        "        |\n"
        "        v\n"
        " MetricsBus -> JSON / matplotlib / 甘特图 / overlap 图",
    )
    add_p(
        doc,
        "一次 decode iteration 的标准时序：时钟弹出 DecodeTick → 调度器选出 batch → 配额与 MPS 确认本片可运行租户 → 对每个请求推进一个 token → 页表 append 可能触发 page fault → 分配器向 VMM 申请驻留页 → 预取运行时并行发起下一层权重搬运并累计 stall → 完成者离队、空槽立即填入就绪队列（持续批）。",
    )

    add_heading_cn(doc, "3.2  模块职责与禁止越权", 2)
    add_table(
        doc,
        ["模块", "允许做", "禁止做"],
        [
            ["DiscreteEventSim", "维护 now 与事件堆、保证因果序", "包含任何调度或分配策略"],
            ["Scheduler", "入队、选批、抢占、统计 TTFT/吞吐", "直接操作物理页或 NVMe 带宽"],
            ["QuotaManager", "份额、硬上限、OOM 拒绝", "决定谁先 decode"],
            ["PrefetchRuntime", "计算/搬运重叠与 stall 度量", "修改页表映射"],
            ["GPUVMM", "驻留、换出、balloon、越界拦截", "理解 token 或 Attention"],
            ["PageTable/Allocator", "逻辑序列↔KV 块、COW、碎片指标", "实现租户时间片"],
        ],
    )

    add_heading_cn(doc, "3.3  仿真与真实引擎解耦", 2)
    add_p(
        doc,
        "系统分为三层适配：Workload 层只产出请求轨迹；Kernel 层（本方案主体）只消费抽象资源代价（step_cost、带宽、页大小）；Adapter 层可选地把同一轨迹接到 HuggingFace transformers 或 vLLM，用于核对“分页降低 OOM / 持续批提高吞吐”的方向是否与仿真一致。Adapter 失败不得阻塞 Kernel 交付。",
    )
    add_table(
        doc,
        ["层", "推荐技术", "成本"],
        [
            ["Kernel / 仿真", "Python ≥3.10，标准库 heapq + asyncio；numpy 仅用于统计", "¥0"],
            ["测试", "pytest、coverage", "¥0"],
            ["可视化", "matplotlib；可选 plotly/gradio 演示页", "¥0"],
            ["真实 GPU", "Colab T4 / Kaggle / AutoDL 按量", "¥0–5/小时"],
            ["推理对照", "transformers 连续 KV 基线；可选 vLLM 参考", "开源"],
            ["数据集", "ShareGPT、LongBench 风格合成、Azure LLM trace、泊松流", "开源"],
        ],
    )

    add_heading_cn(doc, "3.4  仓库结构", 2)
    add_code(
        doc,
        "llmOS/\n"
        "  pyproject.toml                 # 依赖与 pytest 配置\n"
        "  README.md                      # 一键：pytest && python -m llmos.scripts.run_ab\n"
        "  configs/default.yaml           # 冻结带宽、block_size、batch_size、seed\n"
        "  llmos/\n"
        "    sim/clock.py                 # DiscreteEventSim\n"
        "    sim/metrics.py               # 统一指标总线\n"
        "    task1_kv_memory/\n"
        "      kv_block_allocator.py\n"
        "      page_table.py\n"
        "    task2_scheduler/\n"
        "      request.py\n"
        "      events.py\n"
        "      scheduler.py\n"
        "    task3_runtime/\n"
        "      storage_model.py\n"
        "      prefetch_runtime.py\n"
        "      quota_manager.py\n"
        "    task4_gpu_virt/\n"
        "      gpu_vmm.py\n"
        "      mps_scheduler.py\n"
        "    workloads/\n"
        "      synthetic.py\n"
        "      trace_loader.py\n"
        "    adapters/                    # 可选云端对照\n"
        "  tests/\n"
        "  scripts/run_ab.py\n"
        "  docs/设计文档.md  评测报告.md",
    )

    add_heading_cn(doc, "4  子系统详细设计", 1)
    add_heading_cn(doc, "4.1  任务一：KV Cache 分页内存管理", 2)
    add_heading_cn(doc, "4.1.1  基线与优化定义", 3)
    add_p(
        doc,
        "基线采用连续预占：序列创建时一次性申请 ceil(max_tokens / block_size) 个块并在物理上视为必须连续的大段。该模型同时制造内部碎片（实际生成远小于上限）与外部碎片（多序列争抢连续大洞），并无法在 beam / 并行采样间共享前缀。",
    )
    add_p(
        doc,
        "优化采用 PagedAttention 同构设计：物理层是固定大小的 KV block 池（默认 block_size=16 个 token）；每个序列维护页表，将逻辑 token 下标映射为 (block_id, offset)。生成过程按需分配；fork 时只增加引用计数；写路径在 ref_count>1 时拷贝分裂；池耗尽时按 last_access 做 LRU 换出到 HostSwap（CPU 内存模拟区）。",
    )

    add_heading_cn(doc, "4.1.2  数据结构", 3)
    add_code(
        doc,
        "@dataclass\n"
        "class Block:\n"
        "    block_id: int\n"
        "    ref_count: int = 0\n"
        "    is_shared: bool = False\n"
        "    last_access: float = 0.0\n"
        "    resident: bool = True          # False 表示已换出到 HostSwap\n"
        "    swap_location: Optional[int] = None\n"
        "\n"
        "class BlockAllocator:\n"
        "    # free_bitmap: 空闲块位图，O(1) 扫描可用 next-fit / 空闲栈\n"
        "    # blocks: List[Block]\n"
        "    # host_swap: Dict[int, bytes_like]  # 仅记账容量，不存真实张量\n"
        "\n"
        "class PageTable:\n"
        "    # seqs: Dict[seq_id, List[block_id]]  逻辑块号 -> 物理块\n"
        "    # token_len: Dict[seq_id, int]",
    )
    add_p(
        doc,
        "分配器第一期使用“空闲栈 + 引用计数”，保证 allocate/free 摊还 O(1)，满足单测与课程进度。伙伴系统作为加分项：仅当需要演示外部碎片的经典曲线时启用，不作为联调阻塞项。",
    )

    add_heading_cn(doc, "4.1.3  关键算法", 3)
    add_p(doc, "（1）按需扩页 append_token(seq_id)")
    add_bullet(doc, "若 token_len % block_size != 0：只增加逻辑长度，写当前块 offset。")
    add_bullet(doc, "否则视为 page fault：allocate(1)。若空闲不足，steal_for_eviction(now) 选择 ref_count==1 且 last_access 最小的驻留块换出。")
    add_bullet(doc, "换出后仍不足则抛 AllocError，由调度器将该请求标记失败（计入 OOM 比例）。")
    add_p(doc, "（2）写时复制 fork(src_seq)")
    add_bullet(doc, "新序列页表复制源序列的 block_id 列表，对每个块 ref_count+=1，is_shared=True。")
    add_bullet(doc, "任一方再 append 且将写满新位置时，若目标块 is_shared：分配新块、拷贝元数据、源块 ref_count-=1。")
    add_p(doc, "（3）指标")
    add_bullet(doc, "utilization = 实际占用 token 当量 / 物理池 token 当量。")
    add_bullet(doc, "fragmentation_ratio = 1 - 最大可连续满足的 allocate(n) 能力（位图实现下近似为空闲块未能被逻辑序列使用的比例）；内部碎片单独统计 last-block 空洞。")
    add_bullet(doc, "COW 收益 = 共享块数 / 若无共享应占块数。")

    add_heading_cn(doc, "4.1.4  与 VMM 的衔接", 3)
    add_p(
        doc,
        "BlockAllocator.allocate 在联调模式下不直接认为“有空闲下标即成功”，而是对每个新 block 调用 GPUVMM.access(tenant, vpn, write=True)。VMM 负责物理页驻留与超卖换出。这样任务一的 OOM 与任务四的超卖成为同一资源池上的两种视图，避免两套账。单测阶段允许用 FakeVMM（恒驻留）隔离任务一。",
    )

    add_heading_cn(doc, "4.2  任务二：持续批处理请求调度器", 2)
    add_heading_cn(doc, "4.2.1  请求 PCB 与两阶段代价模型", 3)
    add_p(
        doc,
        "Request 即进程控制块。状态：WAITING（就绪）、RUNNING（本 iteration 在批内）、PREEMPTED（被踢出批但 KV 仍在分页器）、FINISHED。调度指标全部从 PCB 时间戳推导，禁止事后用全局计数器“估一个吞吐”。",
    )
    add_p(
        doc,
        "Prefill 视为一次长 CPU burst：cost_prefill = alpha * prompt_len。Decode 每步视为短 burst：cost_decode = beta * f(batch_size)，默认 f 为线性或弱超线性，以体现 Attention 随 batch 的开销。上下文切换开销 cost_cs 只在请求进出 batch 时计入。正因为 KV 常驻（或由 VMM 决定是否换出），iteration 级调度才划算——这是必须写进设计文档的因果说明。",
    )

    add_heading_cn(doc, "4.2.2  静态批 vs 持续批", 3)
    add_p(
        doc,
        "StaticBatchingScheduler：凑满 batch_size 才启动；批内任一请求未结束则整批占用 GPU；短请求完成后槽位空转。这是 OS 早期“不可抢占批处理”的对应物，作为全部对比的基线。",
    )
    add_p(
        doc,
        "ContinuousBatchingScheduler：每个 DecodeTick 重新 _pick_batch。完成者立即释放槽位，WAITING 队头填入。GPU 忙闲比定义为：有非空 batch 的时间 / 仿真时长。",
    )

    add_heading_cn(doc, "4.2.3  策略实现规格", 3)
    add_table(
        doc,
        ["Strategy", "选批规则", "抢占", "预期教学现象"],
        [
            ["FCFS", "按 arrival_time 填满槽位", "否", "吞吐高于静态批；长请求可占槽"],
            ["SJF", "按剩余生成长度升序", "否", "平均周转下降；长请求饥饿"],
            ["PREEMPTIVE", "优先级；时间片用尽踢到 PREEMPTED", "是", "TTFT 改善；需计 ctx 开销"],
            ["MLFQ", "用满片降级，让出升入高队列", "是", "兼顾响应与吞吐（加分）"],
            ["FAIR", "租户令牌桶 / 份额加权", "可", "P99 方差下降，防跨租户霸占"],
        ],
    )
    add_p(
        doc,
        "饥饿定义：等待时间超过阈值 θ（配置项，默认 10× 平均 decode 时长）仍未进入 batch 的请求占比。FAIR 必须把该指标压到配置上限以下，否则视为公平策略未完成。",
    )

    add_heading_cn(doc, "4.2.4  调度主循环", 3)
    add_code(
        doc,
        "def _maybe_schedule(self):\n"
        "    self._reap_finished()\n"
        "    batch = self._pick_batch()          # 受 QuotaManager / MPS 约束\n"
        "    if not batch:\n"
        "        self.sim.schedule(self.idle_tick, self._maybe_schedule)\n"
        "        return\n"
        "    cost = self._iteration_cost(batch)  # prefill/decode 混合\n"
        "    def on_tick():\n"
        "        now = self.sim.now\n"
        "        for req in batch:\n"
        "            req.step(now)               # 维护 first_token_time\n"
        "            self.page_table.append_token(req.seq_id)\n"
        "        self._maybe_schedule()\n"
        "    self.sim.schedule(cost, on_tick)",
    )

    add_heading_cn(doc, "4.3  任务三：权重预取、重叠与多租户配额", 2)
    add_heading_cn(doc, "4.3.1  带宽悬崖模型", 3)
    add_p(
        doc,
        "BandwidthModel 固化三级带宽：NVMe 3.5 GB/s、CPU↔GPU(PCIe) 16 GB/s、GPU HBM 800 GB/s。transfer_time(size, src, dst) 取路径上的瓶颈带宽。禁止在实验中途修改这些常数；若做灵敏度分析，必须单独成节并列出配置哈希。",
    )
    add_p(
        doc,
        "层权重以列表 layer_weights_gb 输入。默认构造一个“前层较小、后层均匀”的合成模型，使预取收益可见，又不必绑定真实 70B 参数文件。",
    )

    add_heading_cn(doc, "4.3.2  双缓冲与 stall 定义", 3)
    add_p(
        doc,
        "PrefetchRuntime 对 decode_step(layer) 并发执行：计算当前层（模拟 GPU 忙碌）与预取 layer+1 … layer+lookahead。I/O stall 定义为 max(0, prefetch_remaining - compute_time)。io_stall_ratio = 累计 stall / 累计（compute + stall）。串行基线 lookahead=0，即算完再搬。",
    )
    add_p(
        doc,
        "权重页缓存采用 LRU，容量为 cache_layers。命中则搬运时间为 0（已在 GPU）。ARC 作为加分淘汰策略。缓存与 VMM 的关系：缓存占用计入租户显存配额，alloc_mem 失败则该层必须走 NVMe 路径或拒绝运行。",
    )

    add_heading_cn(doc, "4.3.3  类 cgroups 配额", 3)
    add_p(
        doc,
        "TenantQuota 含 gpu_compute_share（相对权重）与 gpu_mem_gb（硬上限）。QuotaManager.alloc_mem 超限返回 False，由调用方转为 OOM，不得静默借用其他租户。compute_slice(tenant, tick) 按份额分配本 tick 可用 GPU 时间。is_isolated 自检：任何租户累计显存不超过自身上限，且全局之和不超过物理容量。",
    )
    add_p(
        doc,
        "调度器在 _pick_batch 中必须查询配额：份额用尽的租户本 tick 不能被选入 batch。由此任务二与任务三在控制面上闭合。",
    )

    add_heading_cn(doc, "4.4  任务四：GPU 虚拟化与显存超卖", 2)
    add_heading_cn(doc, "4.4.1  地址空间与页错误", 3)
    add_p(
        doc,
        "每个租户拥有独立虚拟页表。create_tenant(tenant, vmem_pages) 允许 vmem_pages > phys_pages，从而形成超卖。access(tenant, vpn, write, now) 流程：",
    )
    add_bullet(doc, "vpn 越界：记 isolation_violation 尝试次数，拦截并抛 IsolationError，不得分配。")
    add_bullet(doc, "未驻留：page fault；若有空闲 PPN 则建立映射；否则 LRU 换出其他租户冷页到主机 swap。")
    add_bullet(doc, "写访问置 dirty；换出脏页计入额外 swap_write 时间。")
    add_bullet(doc, "oversubscription_ratio = 各租户虚拟页之和 / 物理页数。")
    add_bullet(doc, "isolation_violation_rate 定义为未拦截次数 / 越界尝试次数，验收标准为 0。")

    add_heading_cn(doc, "4.4.2  Balloon 与 MPS", 3)
    add_p(
        doc,
        "balloon_reclaim(tenant, pages) 从该租户选择未访问或可丢弃的空闲虚拟页，解除驻留并归还物理池。它模拟客户机主动交还内存，是超卖压力下的主动回收，而不是错误路径。",
    )
    add_p(
        doc,
        "MPSScheduler 按 share 加权轮转时间片。run_slice 扣除 ctx_switch_cost（默认 0.5 ms 量级，可配置）。fairness_index 采用 Jain 指数：令 x_i 为租户 i 的归一化服务时间（实际服务 / 份额），J = (Σx)^2 / (n Σx^2)，越接近 1 越公平。",
    )

    add_heading_cn(doc, "5  接口契约（实现冻结）", 1)
    add_p(
        doc,
        "以下签名与课程骨架对齐。并行开发期间禁止改名、改参数顺序；需要扩展时以可选参数或新方法追加，并由架构师更新本节。",
    )
    add_heading_cn(doc, "5.1  任务一", 2)
    add_code(
        doc,
        "class BlockAllocator:\n"
        "    def __init__(self, num_blocks: int, block_size: int): ...\n"
        "    def allocate(self, n: int) -> List[int]: ...\n"
        "    def free(self, block_ids: List[int]) -> None: ...\n"
        "    def fork(self, src: int) -> int: ...\n"
        "    def steal_for_eviction(self, now: float) -> Optional[int]: ...\n"
        "    def fragmentation_ratio(self) -> float: ...\n"
        "    def utilization(self) -> float: ...\n"
        "\n"
        "class PageTable:\n"
        "    def __init__(self, allocator: BlockAllocator): ...\n"
        "    def new_sequence(self) -> int: ...\n"
        "    def append_token(self, seq_id: int) -> None: ...\n"
        "    def translate(self, seq_id: int, logical_idx: int) -> tuple[int, int]: ...\n"
        "    def fork(self, src_seq: int) -> int: ...\n"
        "    def free(self, seq_id: int) -> None: ...",
    )
    add_heading_cn(doc, "5.2  任务二", 2)
    add_code(
        doc,
        "class ReqState(Enum):\n"
        "    WAITING = \"waiting\"\n"
        "    RUNNING = \"running\"\n"
        "    PREEMPTED = \"preempted\"   # 扩展，不破坏原有三态语义\n"
        "    FINISHED = \"finished\"\n"
        "\n"
        "class DiscreteEventSim:\n"
        "    def now(self) -> float: ...\n"
        "    def schedule(self, delay: float, fn: Callable[[], None], priority: int = 0): ...\n"
        "    def run(self, until: float): ...\n"
        "    def step(self): ...\n"
        "\n"
        "class BaseScheduler(ABC):\n"
        "    def submit(self, req: Request): ...\n"
        "    def _pick_batch(self) -> List[Request]: ...\n"
        "    def metrics(self) -> Dict[str, float]: ...",
    )
    add_heading_cn(doc, "5.3  任务三 / 四", 2)
    add_code(
        doc,
        "class BandwidthModel:\n"
        "    def transfer_time(self, gigabytes: float, src: str, dst: str) -> float: ...\n"
        "\n"
        "class PrefetchRuntime:\n"
        "    async def decode_step(self, layer: int): ...\n"
        "    def io_stall_ratio(self) -> float: ...\n"
        "\n"
        "class QuotaManager:\n"
        "    def register(self, q: TenantQuota): ...\n"
        "    def alloc_mem(self, tenant: str, gb: float) -> bool: ...\n"
        "    def compute_slice(self, tenant: str, tick: float) -> float: ...\n"
        "    def is_isolated(self, tenant: str) -> bool: ...\n"
        "\n"
        "class GPUVMM:\n"
        "    def create_tenant(self, tenant: str, vmem_pages: int): ...\n"
        "    def access(self, tenant: str, vpn: int, write: bool, now: float): ...\n"
        "    def balloon_reclaim(self, tenant: str, pages: int): ...\n"
        "    def oversubscription_ratio(self) -> float: ...\n"
        "    def isolation_violation_rate(self) -> float: ...\n"
        "\n"
        "class MPSScheduler:\n"
        "    def pick_next(self, now: float) -> str: ...\n"
        "    def run_slice(self, tenant: str, now: float) -> float: ...\n"
        "    def fairness_index(self) -> float: ...",
    )
    add_p(
        doc,
        "错误语义：AllocError 表示物理 KV 池与换出后仍不足；IsolationError 表示越界；QuotaOOM 表示硬上限拒绝。调度器必须捕获并计入失败指标，不得让仿真时钟崩溃。",
    )

    add_heading_cn(doc, "6  评测体系", 1)
    add_heading_cn(doc, "6.1  工作负载", 2)
    add_table(
        doc,
        ["负载档", "构造方法", "考察点"],
        [
            ["W1 短对话", "ShareGPT 截断或合成 prompt 32–128、生成 32–128", "持续批与 TTFT"],
            ["W2 长上下文", "合成 prompt 2k–8k、生成 256–1024", "KV 碎片、OOM、换出"],
            ["W3 多租户突发", "两租户泊松到达，份额 1:1 与 3:1", "FAIR、配额、MPS 公平"],
            ["W4 并行采样", "同一 prompt fork 4 路", "COW 共享收益"],
        ],
    )
    add_p(
        doc,
        "所有 A/B 实验写入 configs 快照：seed、到达率、block_size、batch_size、带宽、时间片。scripts/run_ab.py 输出 results/<exp_id>.json 与对应 PNG，评测报告只引用这些文件，禁止手抄数字。",
    )

    add_heading_cn(doc, "6.2  实验矩阵", 2)
    add_table(
        doc,
        ["编号", "对比", "主指标", "通过线（设计预期）"],
        [
            ["E1", "连续预占 vs 分页", "utilization、fragmentation、并发序列、OOM 率", "碎片显著下降，OOM 不升反降"],
            ["E2", "分页 vs 分页+COW", "同等采样下块占用", "W4 下块数接近 1/分支数 + 增量"],
            ["E3", "分页 vs 分页+LRU 换出", "OOM 率、换入延迟", "压力场景 OOM 下降，延迟可解释"],
            ["E4", "静态批 vs 持续批 FCFS", "req/s、TTFT P50/P99、GPU 忙闲比", "吞吐与忙闲比上升"],
            ["E5", "FCFS vs SJF vs 抢占", "周转、饥饿占比", "SJF 饥饿可见；抢占改善 TTFT"],
            ["E6", "FCFS vs FAIR（W3）", "各租户 P99 方差、饥饿", "方差下降，无长期饥饿"],
            ["E7", "串行 vs 双缓冲 vs 缓存", "io_stall_ratio、有效算力", "stall 单调下降"],
            ["E8", "无配额 vs 硬上限", "隔离自检、他租户延迟", "越界申请被拒，邻租户稳定"],
            ["E9", "无超卖 vs 超卖+balloon", "超卖比、延迟、拦截率", "拦截率=0，延迟随超卖比上升"],
            ["E10", "MPS 份额 1:1 vs 3:1", "Jain 指数、服务时间比", "服务时间近似跟份额"],
        ],
    )

    add_heading_cn(doc, "6.3  云端对照（非阻塞）", 2)
    add_p(
        doc,
        "只做一项主对照即可答辩：优先用 transformers 在 T4 上复现“连续 KV 预分配导致并发下降/OOM”，与 E1 方向对照。若环境允许，再用 vLLM 观察持续批相对静态批的吞吐方向，对照 E4。记录 nvidia-smi、模型名、精度、脚本版本。对照失败时在报告中写明环境限制，仍以仿真为准。",
    )

    add_heading_cn(doc, "6.4  可视化清单", 2)
    add_bullet(doc, "调度甘特图：请求占用 batch 槽位随时间变化（静态批空泡一目了然）。")
    add_bullet(doc, "KV 利用率与碎片时间序列。")
    add_bullet(doc, "计算与预取 overlap 图（双缓冲交错）。")
    add_bullet(doc, "超卖比–P99 延迟曲线。")
    add_bullet(doc, "多租户 P99 柱状图（FAIR 前后）。")

    add_heading_cn(doc, "7  工程实现方案", 1)
    add_heading_cn(doc, "7.1  技术选型", 2)
    add_p(
        doc,
        "时钟用标准库 heapq 实现优先队列，不强制依赖 simpy，降低环境摩擦；PrefetchRuntime 用 asyncio 表达并发搬运，在离散事件中把 await 点翻译为 schedule(delay, callback)，避免真实睡眠。统计可用纯 Python，热点循环若需要再用 numpy。测试用 pytest；覆盖率用 pytest-cov。配置用 PyYAML。",
    )
    add_heading_cn(doc, "7.2  质量门禁", 2)
    add_bullet(doc, "每次 PR：pytest -q --cov=llmos --cov-fail-under=70。")
    add_bullet(doc, "黄金 trace（固定 seed 的 W1 子集）指标波动超过 1% 视为回归。")
    add_bullet(doc, "角色 B 改分配器必须由 C review；C 改调度必须由 B review。")
    add_bullet(doc, "禁止提交真实 API Key、云端账号与未脱敏日志。")

    add_heading_cn(doc, "7.3  确定性", 2)
    add_p(
        doc,
        "离散事件不得依赖线程调度顺序。所有随机只来自 random.Random(seed)。asyncio 路径必须把完成时间写入事件队列，而不是依赖 Task 完成先后。这样同一配置在任意机器上数字一致，满足“可复现”考核。",
    )

    add_heading_cn(doc, "8  组织、进度与实习节点映射", 1)
    add_heading_cn(doc, "8.1  角色", 2)
    add_table(
        doc,
        ["角色", "代号", "主交付", "技能"],
        [
            ["系统架构师", "A", "概念映射、模块/时序图、接口冻结、进度、设计文档统稿", "系统思维、UML"],
            ["内存/系统实现", "B", "页表、分配器、配额、GPUVMM、相关单测", "数据结构、并发"],
            ["运行时/IO 实现", "C", "时钟、调度策略、预取流水线、MPS、引擎对接", "asyncio、事件仿真"],
            ["测试与评测", "D", "负载、A/B 脚本、图表、覆盖率、评测报告", "pytest、pandas、matplotlib"],
        ],
    )

    add_heading_cn(doc, "8.2  十二周计划（每周 8 学时）", 2)
    add_table(
        doc,
        ["周", "里程碑", "关键产出", "实习节点"],
        [
            ["1", "OS 映射 + 仓库骨架 + 时钟可跑", "概念图、空接口、Hello Sim", "环境熟悉"],
            ["2", "接口冻结 + 单测框架 + 配置体系", "设计文档 v1", "系统设计"],
            ["3", "位图/栈式分配器 + 页表增删 token", "任务一单测绿", "系统实现"],
            ["4", "静态批基线 + 持续批 FCFS", "E4 初值", "系统实现"],
            ["5", "变长生成 + Prefill/Decode 代价分离", "W1 可跑通", "系统实现"],
            ["6", "COW + 碎片对比仿真", "E1/E2 图表", "系统实现"],
            ["7", "LRU 换出 + 抢占时间片 + 串行 I/O 基线", "E3/E5、stall 基线", "系统实现"],
            ["8", "双缓冲预取 + 配额接入选批", "E7/E8", "系统实现"],
            ["9", "GPUVMM 按需分页 + MPS", "E9/E10 初值", "系统实现"],
            ["10", "balloon + 全链路联调 + 可选云端对照", "黄金 trace 稳定", "系统测试"],
            ["11", "覆盖率、可视化、文档", "评测报告 v1、README", "系统测试"],
            ["12", "答辩、交叉评审、录屏", "PPT + 5 分钟演示", "演示汇报"],
        ],
    )
    add_p(
        doc,
        "若实习将“系统设计答辩通过后才给设备经费”，本方案主动把 GPU 依赖后置：前 9 周零硬件可完成全部逻辑与评测草稿，不阻塞进度。",
    )

    add_heading_cn(doc, "8.3  过程考核对齐", 2)
    add_bullet(doc, "过程 40%：周会记录、提交活跃度、B/C 互审记录。")
    add_bullet(doc, "结果 60%：设计文档映射清晰度 20%；实现与可复现 20%；量化对比 15%；答辩 5%。")
    add_bullet(doc, "实习评分另含场景覆盖、新颖性、功能完成度、新技术、文档、参与度——本方案用“四子系统合一的 llmOS”覆盖 OS 知识点与新颖性，用冻结实验矩阵覆盖功能完成度。")

    add_heading_cn(doc, "9  风险、砍法与伦理", 1)
    add_table(
        doc,
        ["风险", "影响", "应对"],
        [
            ["四模块同时做深做不完", "联调失败", "先通一条请求全路径；MLFQ、ARC、伙伴、vLLM 对接标加分"],
            ["仿真参数调出假提升", "结论不可信", "冻结 default.yaml；灵敏度单独成节"],
            ["云端 GPU 不稳定", "对照缺失", "仿真为主交付，对照进附录"],
            ["接口漂移", "集成爆炸", "签名冻结 + 互审 + 契约测试"],
            ["事件仿真写成就绪循环", "结果依赖机器", "禁止真实 sleep；完成时间入队"],
            ["数据与账号泄漏", "安全事故", "不用生产密钥；trace 脱敏"],
        ],
    )
    add_p(
        doc,
        "最低可答辩集合（必须同时具备，否则不开答辩）：E1、E4、E7、E9 四组对比图 + 覆盖率达标 + 一键复现脚本。其余实验作为消融与加分。",
    )

    add_heading_cn(doc, "10  最终交付清单", 1)
    add_bullet(doc, "Git 代码仓：含 README、环境安装、一键仿真与评测。")
    add_bullet(doc, "本文《完整技术方案》及后续修订版（接口变更必须改文档）。")
    add_bullet(doc, "《系统设计文档》附图：OS 映射图、模块图、decode iteration 时序图、页错误时序图。")
    add_bullet(doc, "《量化评测报告》：E1–E10 能做多少做多少，但 E1/E4/E7/E9 为必交。")
    add_bullet(doc, "答辩 PPT + 5 分钟演示录屏（现场跑黄金 trace）。")
    add_bullet(doc, "过程材料：周报、PR 互审记录。")

    add_heading_cn(doc, "11  实施要点（给实现者的硬约束）", 1)
    add_p(doc, "下列条目写入仓库 CONTRIBUTING，作为代码审查清单：")
    add_bullet(doc, "PageTable.translate 必须对非法 seq_id / 越界 logical_idx 失败，不得返回随机块。")
    add_bullet(doc, "free 只有在 ref_count 降到 0 时才把块放回空闲栈，防止 COW 下提前回收。")
    add_bullet(doc, "steal_for_eviction 不得驱逐 ref_count>1 的共享块，除非先分裂。")
    add_bullet(doc, "Static 与 Continuous 必须共享同一 Request.step 与同一代价函数，对比才公平。")
    add_bullet(doc, "FAIR 的令牌发放与 MPS 的份额必须用同一 tenant 配置源，避免两套份额。")
    add_bullet(doc, "Prefetch 的 stall 只统计“计算线程因数据未就绪而等待”的时间，不得把 DMA 忙碌算进 GPU 有效算力。")
    add_bullet(doc, "GPUVMM.isolation_violation_rate 在故意越界测试中分母>0、分子必须为 0。")
    add_bullet(doc, "run_ab 默认打印基线与优化两列，缺任一列视为评测脚本不合格。")

    add_heading_cn(doc, "12  结论", 1)
    add_p(
        doc,
        "本技术方案将课程给出的四个操作系统级任务收敛为同一套 llmOS 内核：用离散事件仿真复现分页内存、持续批调度、计算-I/O 重叠与 GPU 超卖，并用冻结的接口、实验矩阵和质量门禁保证可复现、可分工、可答辩。原理仍是操作系统课本中的页表、抢占、双缓冲与内存超卖；变化的只是 workload——从进程变成了推理请求与 KV Cache。该路线同时满足 A 类选题对系统深度的要求，以及“云端 GPU + 本地仿真、无需本地硬件”的排课约束。",
    )

    add_heading_cn(doc, "附录 A  默认配置（冻结稿）", 1)
    add_code(
        doc,
        "seed: 42\n"
        "block_size: 16\n"
        "num_kv_blocks: 1024\n"
        "batch_size: 8\n"
        "alpha_prefill: 0.00002    # 秒/token\n"
        "beta_decode: 0.001        # 秒/iteration·batch\n"
        "nvme_gbps: 3.5\n"
        "cpu_gpu_gbps: 16.0\n"
        "gpu_hbm_gbps: 800.0\n"
        "lookahead: 2\n"
        "time_slice: 0.01\n"
        "ctx_switch_cost: 0.0005\n"
        "phys_pages: 64\n"
        "page_size_gb: 0.25\n"
        "hunger_threshold_x: 10",
    )

    add_heading_cn(doc, "附录 B  设计评审检查表（第 2 周答辩用）", 1)
    add_bullet(doc, "是否能在一页纸上画出请求穿过四个子系统的路径？")
    add_bullet(doc, "是否写清静态批与持续批的差异，以及为何 KV 常驻使细粒度调度可行？")
    add_bullet(doc, "是否定义碎片、stall、超卖比、Jain 指数的计算公式？")
    add_bullet(doc, "是否列出 E1/E4/E7/E9 为必做，其余为加分？")
    add_bullet(doc, "是否说明无 GPU 也能完成主交付？")
    add_bullet(doc, "B/C 接口是否已冻结并有契约测试清单？")

    doc.save(OUT)
    print("WROTE", OUT)


if __name__ == "__main__":
    build()
